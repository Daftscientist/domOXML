#!/usr/bin/env python3
# pyright: basic
# A dev tool built on python-pptx and lxml, which are only partly typed; ruff still applies.
"""Native-recipe lab: ask LibreOffice and real PowerPoint what an OOXML feature actually does.

Raster fallbacks exist for CSS features with no obvious DrawingML equivalent. Before building a
native recipe, probe the candidate here: each probe writes a small deck exercising one feature,
renders it in LibreOffice and (if the local Graph tool is set up) real PowerPoint, and lays the
two renders side by side. Reading the sheet tells you what is honoured, what is ignored and what
Office rejects, which is what the recipes in ``docs/native-recipes.md`` were built from.

    python scripts/native_lab.py colour scene3d gradients blur-edges blur-scale
    python scripts/native_lab.py gradients --no-graph        # LibreOffice only

Real PowerPoint needs ``.local/graph_render.py`` and ``source .local/graph.env`` first. Office
rejects markup it considers invalid with ``PPTCannotOpenFileMetroFileCorruption``; the lab prints
that error code, which doubles as a free schema check. Text metrics in Graph renders are not
comparable (embedded fonts are stripped), so probe shapes, fills and effects, not wrapping.
"""

from __future__ import annotations

import argparse
import base64
import colorsys
import importlib
import json
import re
import subprocess
import sys
import urllib.error
import urllib.request
import uuid
from collections.abc import Callable
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide
from pptx.util import Emu, Inches, Pt

# lxml ships no type stubs; the repo keeps that behind one dynamic import (see core/capabilities).
etree: Any = importlib.import_module("lxml.etree")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "out" / "native-lab"
GRAPH_TOOL = ROOT / ".local" / "graph_render.py"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

type Build = Callable[[PresentationType, Slide], None]


def _q(namespace: str, tag: str) -> str:
    return f"{{{namespace}}}{tag}"


def _xml(markup: str) -> Any:
    return etree.fromstring(f'<r xmlns:a="{A}" xmlns:p="{P}">{markup}</r>')[0]


def photo() -> bytes:
    """A colourful test image: a hue sweep, three squares and a black bar."""
    image = Image.new("RGB", (400, 220))
    pixels = image.load()
    assert pixels is not None
    for x in range(400):
        red, green, blue = colorsys.hsv_to_rgb(x / 400, 0.75, 0.95)
        for y in range(220):
            pixels[x, y] = (int(red * 255), int(green * 255), int(blue * 255))
    draw = ImageDraw.Draw(image)
    draw.rectangle((30, 30, 130, 120), fill=(230, 40, 40))
    draw.rectangle((150, 60, 250, 150), fill=(40, 160, 60))
    draw.rectangle((270, 30, 370, 120), fill=(40, 70, 220))
    draw.rectangle((0, 170, 400, 190), fill="black")
    return _png(image)


def stripes() -> bytes:
    """Hard black and white stripes with a red box: blur and clipping show up clearly."""
    image = Image.new("RGB", (600, 300), "white")
    draw = ImageDraw.Draw(image)
    for x in range(0, 600, 40):
        draw.rectangle((x, 0, x + 19, 300), fill="black")
    draw.rectangle((220, 90, 380, 210), fill=(220, 30, 30))
    return _png(image)


def _png(image: Image.Image) -> bytes:
    buffer = BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def _deck(name: str, build: Build, width: float = 13.333, height: float = 7.5) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(width), Inches(height)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    build(presentation, slide)
    path = OUT / f"{name}.pptx"
    presentation.save(str(path))
    return path


def _caption(slide: Slide, text: str, x: float, y: float, width: float = 3.0) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(0.3))
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(11)


# --------------------------------------------------------------------------- rendering


def _graph_module():
    sys.path.insert(0, str(GRAPH_TOOL.parent))
    import graph_render  # type: ignore[import-not-found]

    return graph_render


def office_error(path: Path) -> str:
    """Upload ``path`` to real PowerPoint and return Office's error code, or ``"ok"``."""
    graph = _graph_module()
    token = graph._token(interactive=False)
    data = graph._strip_embedded_fonts(path.read_bytes())
    name = f"domoxml-tmp-{uuid.uuid4().hex}.pptx"
    uploaded = graph._request(
        "PUT",
        f"{graph.GRAPH}/me/drive/root:/{name}:/content",
        token,
        data=data,
        ctype=graph._PPTX_CT,
    )
    item = json.load(uploaded)["id"]
    try:
        request = urllib.request.Request(
            f"{graph.GRAPH}/me/drive/items/{item}/content?format=pdf",
            headers={"Authorization": f"Bearer {token}"},
        )
        try:
            urllib.request.urlopen(request)
        except urllib.error.HTTPError as error:
            match = re.search(r"ErrorCode=(\w+)", error.read().decode())
            return match.group(1) if match else f"HTTP {error.code}"
        return "ok"
    finally:
        graph._request("DELETE", f"{graph.GRAPH}/me/drive/items/{item}", token)


def render(path: Path, *, graph: bool) -> list[Path]:
    """Render ``path`` in LibreOffice and optionally real PowerPoint; return the PNGs made."""
    made: list[Path] = []
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(path)],
        capture_output=True,
        check=False,
    )
    pngs = {"lo": OUT / f"{path.stem}.pdf"}
    if graph and GRAPH_TOOL.exists():
        target = OUT / f"gp-{path.stem}.pdf"
        subprocess.run(
            [sys.executable, str(GRAPH_TOOL), "render", str(path), str(target)],
            capture_output=True,
            check=False,
        )
        if target.exists():
            pngs["gp"] = target
        else:
            print(f"  real PowerPoint refused {path.name}: {office_error(path)}")
    for tag, pdf in pngs.items():
        stem = OUT / f"{tag}-{path.stem}"
        subprocess.run(
            ["pdftoppm", "-png", "-r", "70", "-singlefile", str(pdf), str(stem)], check=False
        )
        made.append(stem.with_suffix(".png"))
    return made


def sheet(name: str, pngs: list[Path]) -> Path:
    """Lay renders side by side (LibreOffice left, real PowerPoint right)."""
    images = [Image.open(p).convert("RGB") for p in pngs if p.exists()]
    out = Image.new(
        "RGB",
        (sum(i.width for i in images) + 10 * len(images), max(i.height for i in images)),
        "white",
    )
    x = 0
    for image in images:
        out.paste(image, (x, 0))
        x += image.width + 10
    path = OUT / f"sheet-{name}.png"
    out.save(path)
    return path


# --------------------------------------------------------------------------- probes


def probe_colour() -> Path:
    """Which ``a:blip`` colour operations are honoured (candidates for CSS ``filter`` on images)."""
    image = photo()
    variants = [
        ("plain", None),
        ("grayscl", "<a:grayscl/>"),
        ("lum bright+30", '<a:lum bright="30000" contrast="0"/>'),
        ("lum contrast+50", '<a:lum bright="0" contrast="50000"/>'),
        ("biLevel 50", '<a:biLevel thresh="50000"/>'),
        ("duotone", '<a:duotone><a:srgbClr val="1A237E"/><a:srgbClr val="FFEB3B"/></a:duotone>'),
        ("alphaModFix 50%", '<a:alphaModFix amt="50000"/>'),
        (
            "clrChange red->white",
            '<a:clrChange><a:clrFrom><a:srgbClr val="E62828"/></a:clrFrom>'
            '<a:clrTo><a:srgbClr val="FFFFFF"/></a:clrTo></a:clrChange>',
        ),
    ]

    def build(_: PresentationType, slide: Slide) -> None:
        for index, (label, markup) in enumerate(variants):
            x, y = 0.3 + (index % 4) * 3.2, 0.3 + (index // 4) * 2.4
            picture = slide.shapes.add_picture(
                BytesIO(image), Inches(x), Inches(y), Inches(3.0), Inches(1.65)
            )
            if markup:
                picture._element.find(f".//{_q(A, 'blip')}").insert(0, _xml(markup))
            _caption(slide, label, x, y + 1.68)

    return sheet("colour", render(_deck("lab_colour", build), graph=_GRAPH))


def probe_scene3d() -> Path:
    """What 3-D scene presets do to a shape and its text (candidates for flips, perspective)."""
    cases = [
        ("plain", None),
        ("shape lon=180", ("sp", "orthographicFront", 0, 10_800_000)),
        ("shape lat=40", ("sp", "orthographicFront", 2_400_000, 0)),
        ("shape perspectiveLeft", ("sp", "perspectiveLeft", 0, 1_800_000)),
        ("TEXT lon=180 (mirror)", ("tx", "orthographicFront", 0, 10_800_000)),
        ("TEXT lat=40", ("tx", "orthographicFront", 2_400_000, 0)),
        ("shape perspectiveRelaxed", ("sp", "perspectiveRelaxed", 1_800_000, 1_200_000)),
    ]

    def build(_: PresentationType, slide: Slide) -> None:
        for index, (label, spec) in enumerate(cases):
            x, y = 0.4 + (index % 4) * 3.2, 0.4 + (index // 4) * 2.6
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(2.6), Inches(1.5)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x30, 0x60, 0xD0)
            shape.text_frame.text = "Mirror 123"
            if spec:
                where, preset, lat, lon = spec
                scene = _xml(
                    f'<a:scene3d><a:camera prst="{preset}">'
                    f'<a:rot lat="{lat}" lon="{lon}" rev="0"/></a:camera>'
                    '<a:lightRig rig="threePt" dir="t"/></a:scene3d>'
                )
                target = (
                    shape._element.find(_q(P, "spPr"))
                    if where == "sp"
                    else shape._element.find(f".//{_q(A, 'bodyPr')}")
                )
                target.append(scene)
            _caption(slide, label, x, y + 1.55)

    return sheet("scene3d", render(_deck("lab_scene3d", build), graph=_GRAPH))


def _alpha(alpha: int | None) -> str:
    return "" if alpha is None else f'<a:alpha val="{alpha}"/>'


def probe_gradients() -> Path:
    """Many-stop and alpha-stop gradients (repeating gradients and ``mask-image`` fades)."""

    def gradient(stops: list[tuple[int, str, int | None]]) -> Any:
        points = "".join(
            f'<a:gs pos="{pos}"><a:srgbClr val="{colour}">{_alpha(alpha)}</a:srgbClr></a:gs>'
            for pos, colour, alpha in stops
        )
        return _xml(
            f'<a:gradFill rotWithShape="1"><a:gsLst>{points}</a:gsLst>'
            '<a:lin ang="0" scaled="0"/></a:gradFill>'
        )

    hard: list[tuple[int, str, int | None]] = []
    for tile in range(20):
        colour = "000000" if tile % 2 == 0 else "FFFFFF"
        hard += [(tile * 5000, colour, None), ((tile + 1) * 5000 - 1, colour, None)]
    cases = [
        ("20 hard stripes (40 stops)", hard),
        ("alpha fade 100 -> 0", [(0, "1A237E", 100_000), (100_000, "1A237E", 0)]),
        (
            "alpha fade in and out",
            [(0, "D32F2F", 0), (50_000, "D32F2F", 100_000), (100_000, "D32F2F", 0)],
        ),
        (
            "121 smooth stops",
            [
                (
                    round(i * 100_000 / 120),
                    f"{round(255 * i / 120):02X}50{255 - round(255 * i / 120):02X}",
                    None,
                )
                for i in range(121)
            ],
        ),
    ]

    def build(_: PresentationType, slide: Slide) -> None:
        for index, (label, stops) in enumerate(cases):
            x, y = 0.4 + (index % 2) * 6.4, 0.4 + (index // 2) * 2.9
            ground = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(6.0), Inches(2.2)
            )
            ground.fill.solid()
            ground.fill.fore_color.rgb = RGBColor(0xF5, 0xC8, 0x42)
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(x + 0.5), Inches(y + 0.3), Inches(5.0), Inches(1.6)
            )
            properties = shape._element.find(_q(P, "spPr"))
            for old in properties.findall(_q(A, "solidFill")) + properties.findall(
                _q(A, "gradFill")
            ):
                properties.remove(old)
            properties.insert(2, gradient(stops))
            _caption(slide, label, x, y + 2.25, 5.5)

    return sheet("gradients", render(_deck("lab_gradients", build), graph=_GRAPH))


def probe_blur_edges() -> Path:
    """``a:blur`` on a picture, rotated and rounded, with ``grow`` 0 and 1 (backdrop-filter)."""
    image = stripes()
    cases = [
        ("rect, grow=0", 0, "rect", 0),
        ("roundRect, grow=0", 0, "roundRect", 0),
        ("roundRect rot -16, grow=0", -16, "roundRect", 0),
        ("roundRect rot -16, grow=1", -16, "roundRect", 1),
    ]

    def build(_: PresentationType, slide: Slide) -> None:
        for index, (label, rotation, geometry, grow) in enumerate(cases):
            x, y = 0.5 + (index % 2) * 6.4, 0.5 + (index // 2) * 3.6
            slide.shapes.add_picture(BytesIO(image), Inches(x), Inches(y), Inches(6.0), Inches(3.0))
            panel = slide.shapes.add_picture(
                BytesIO(image), Inches(x + 1.2), Inches(y + 0.5), Inches(3.6), Inches(2.0)
            )
            panel.crop_left = panel.crop_right = 0.2
            panel.crop_top = panel.crop_bottom = 0.1667
            panel.rotation = rotation
            properties = panel._element.find(_q(P, "spPr"))
            preset = properties.find(_q(A, "prstGeom"))
            preset.set("prst", geometry)
            if geometry == "roundRect":
                guides = preset.find(_q(A, "avLst"))
                etree.SubElement(
                    guides if guides is not None else etree.SubElement(preset, _q(A, "avLst")),
                    _q(A, "gd"),
                    name="adj",
                    fmla="val 18000",
                )
            effects = etree.SubElement(properties, _q(A, "effectLst"))
            etree.SubElement(effects, _q(A, "blur"), rad="127000", grow=str(grow))
            _caption(slide, label, x, y + 3.05, 5.5)

    return sheet("blur-edges", render(_deck("lab_blur_edges", build, 13.3, 7.8), graph=_GRAPH))


def probe_blur_scale() -> Path:
    """How PowerPoint's blur radius maps to a CSS sigma, scored against Chrome (best scale wins)."""
    from playwright.sync_api import sync_playwright

    image = stripes()
    scales = (1.0, 2.0, 3.0, 4.0)
    width, height = 560, 280
    slots = [(20 + column * 620, 20 + row * 320) for row in range(2) for column in range(2)]
    lines: list[str] = []
    for sigma in (4, 8, 16):
        encoded = base64.b64encode(image).decode()
        panels = "".join(
            f"<div style='position:absolute;left:{x}px;top:{y}px;width:{width}px;"
            f"height:{height}px;overflow:hidden'><img src='data:image/png;base64,{encoded}' "
            f"style='width:{width}px;height:{height}px;filter:blur({sigma}px)'></div>"
            for x, y in slots
        )
        reference = OUT / f"ref-blur-{sigma}.png"
        OUT.mkdir(parents=True, exist_ok=True)
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={"width": 1280, "height": 720})
            page.set_content(f"<body style='margin:0;background:#fff'>{panels}</body>")
            page.screenshot(path=str(reference))
            browser.close()

        def build(_: PresentationType, slide: Slide, sigma: int = sigma) -> None:
            for scale, (x, y) in zip(scales, slots, strict=True):
                picture = slide.shapes.add_picture(
                    BytesIO(image),
                    Emu(x * 9525),
                    Emu(y * 9525),
                    Emu(width * 9525),
                    Emu(height * 9525),
                )
                effects = etree.SubElement(picture._element.find(_q(P, "spPr")), _q(A, "effectLst"))
                etree.SubElement(
                    effects, _q(A, "blur"), rad=str(round(sigma * scale * 9525)), grow="0"
                )

        path = _deck(f"lab_blur_scale_{sigma}", build, 1280 / 96, 720 / 96)
        made = render(path, graph=True)
        candidates = [p for p in made if p.name.startswith("gp-")]
        if not candidates:
            lines.append(f"sigma {sigma}: needs real PowerPoint (set up .local/graph_render.py)")
            continue
        got = Image.open(candidates[0]).convert("L").resize((1280, 720))
        want = Image.open(reference).convert("L")
        inset = 3 * sigma + 8
        scores = [
            (
                scale,
                ImageStat.Stat(
                    ImageChops.difference(
                        want.crop((x + inset, y + inset, x + width - inset, y + height - inset)),
                        got.crop((x + inset, y + inset, x + width - inset, y + height - inset)),
                    )
                ).mean[0],
            )
            for scale, (x, y) in zip(scales, slots, strict=True)
        ]
        best = min(scores, key=lambda s: s[1])[0]
        lines.append(
            f"sigma {sigma:>2}px: "
            + "  ".join(f"x{s:g}={d:5.1f}" for s, d in scores)
            + f"   best x{best:g}"
        )
    report = OUT / "blur-scale.txt"
    report.write_text("\n".join(lines) + "\n")
    print("\n".join(lines))
    return report


PROBES: dict[str, Callable[[], Path]] = {
    "colour": probe_colour,
    "scene3d": probe_scene3d,
    "gradients": probe_gradients,
    "blur-edges": probe_blur_edges,
    "blur-scale": probe_blur_scale,
}
_GRAPH = True


def main() -> None:
    global _GRAPH
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("probes", nargs="+", choices=sorted(PROBES), help="probes to run")
    parser.add_argument("--no-graph", action="store_true", help="LibreOffice only")
    arguments = parser.parse_args()
    _GRAPH = not arguments.no_graph
    for name in arguments.probes:
        print(f"{name}: {PROBES[name]()}")


if __name__ == "__main__":
    main()
