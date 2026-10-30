"""Browserless Takumi runtime smoke tests (requires Node packages installed)."""

from __future__ import annotations

from io import BytesIO

import pytest
from pptx import Presentation as PptxPresentation

from domoxml import Presentation, Slide
from domoxml.types import OutputFormat, Representation

pytestmark = pytest.mark.integration


def test_takumi_html_to_png_and_pptx() -> None:
    deck = Presentation(render_engine="takumi")
    deck.add(
        Slide(
            html=(
                '<div style="position:absolute;left:40px;top:50px;width:220px;'
                'height:80px;background:#4f46e5;color:white;font-size:30px">Hello</div>'
            )
        )
    )
    result = deck.render({OutputFormat.PNG, OutputFormat.PPTX})
    assert len(result.pngs) == 1
    assert result.pngs[0].startswith(b"\x89PNG\r\n\x1a\n")
    assert result.pptx is not None
    assert result.pptx.startswith(b"PK")
    pptx = PptxPresentation(BytesIO(result.pptx))
    assert any(shape.has_text_frame and "Hello" in shape.text for shape in pptx.slides[0].shapes)
    assert result.coverage.items
    assert any(item.representation is Representation.NATIVE for item in result.coverage.items), (
        result.warnings
    )
    assert all(
        item.representation in (Representation.NATIVE, Representation.RASTERIZED)
        for item in result.coverage.items
    )
